"""Excel export — one workbook, one sheet per audience/section."""

from openpyxl import Workbook
from openpyxl.styles import Font


def export_excel(
    path: str,
    period_label: str,
    ticket_kpis: dict,
    network_kpis: dict,
    ai_report: dict,
    technician_performance: dict | None = None,
) -> None:
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Executive Summary"
    summary_ws["A1"] = f"AI IT Operations — Executive Report ({period_label})"
    summary_ws["A1"].font = Font(bold=True, size=14)
    summary_ws["A3"] = "Executive Summary"
    summary_ws["A3"].font = Font(bold=True)
    summary_ws["A4"] = ai_report.get("executive_summary", "Insufficient data")

    row = 6
    for section in ["recurring_problems", "risk_observations", "vendor_observations", "cost_observations", "recommended_actions", "management_attention_items"]:
        summary_ws.cell(row=row, column=1, value=section.replace("_", " ").title()).font = Font(bold=True)
        row += 1
        for item in ai_report.get(section, []):
            summary_ws.cell(row=row, column=1, value=f"- {item}")
            row += 1
        row += 1

    ticket_ws = wb.create_sheet("Ticket KPIs")
    ticket_ws.append(["Metric", "Value"])
    if not ticket_kpis.get("insufficient_data"):
        for key in [
            "total_tickets", "open_tickets", "closed_tickets", "sla_compliance_pct", "mttr_hours",
            "closed_last_7_days", "closed_last_30_days",
            "opened_today_total", "opened_today_assigned", "opened_today_unassigned",
        ]:
            ticket_ws.append([key, ticket_kpis.get(key)])
        ticket_ws.append([])
        ticket_ws.append(["Category", "Count"])
        for cat, count in ticket_kpis.get("category_counts", {}).items():
            ticket_ws.append([cat, count])

    network_ws = wb.create_sheet("Network KPIs")
    network_ws.append(["Site", "Committed Mbps", "Measured Mbps", "Underdelivery %", "Latency ms", "Packet Loss %", "Uptime %", "Provider"])
    if not network_kpis.get("insufficient_data"):
        for site, s in network_kpis.get("site_summaries", {}).items():
            network_ws.append([
                site, s["avg_bandwidth_committed_mbps"], s["avg_bandwidth_measured_mbps"],
                s["bandwidth_underdelivery_pct"], s["avg_latency_ms"], s["avg_packet_loss_pct"],
                s["avg_uptime_pct"], s["circuit_provider"],
            ])

    tech_ws = wb.create_sheet("Technician Performance")
    tech_ws.append(["Rank", "Technician", "Total Tickets", "Closed", "SLA %", "MTTR (hrs)"])
    if technician_performance and not technician_performance.get("insufficient_data"):
        by_technician = technician_performance.get("by_technician", {})
        ranking = technician_performance.get("ranking_best_to_worst", [])
        excluded = set(technician_performance.get("excluded_from_ranking", []))
        ordered = ranking + [t for t in by_technician if t in excluded]
        for rank, tech in enumerate(ordered, start=1):
            perf = by_technician[tech]
            rank_label = rank if tech not in excluded else "excluded"
            tech_ws.append([
                rank_label, tech, perf["total_tickets"], perf["closed_tickets"],
                perf.get("sla_compliance_pct"), perf.get("mttr_hours"),
            ])

    wb.save(path)
